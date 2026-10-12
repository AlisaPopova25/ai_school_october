"""Приблизительный рендер pptx в PNG без PowerPoint: картинки, прямоугольники, текст (Segoe UI вместо Inter), таблицы.

    python poster/preview.py [poster.pptx] [out.png]

Нужен только чтобы проверить вёрстку: наложения, вылезающий текст, баланс колонок. Переносы строк в PowerPoint
могут немного отличаться.
"""
import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "poster.pptx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "preview.png"
PX_PER_CM = 25
INTER_WIDTH = 1.08  # Inter шире Segoe UI (калибровка по экспорту из PowerPoint)
INTER_LINE = 1.245  # межстрочье Inter: 1.245 × 1.1 ≈ 1.37 кегля
EMU_PER_CM = 360000
FONTS = {(False, False): "segoeui.ttf", (True, False): "segoeuib.ttf", (False, True): "segoeuii.ttf",
         (True, True): "segoeuiz.ttf"}
_cache = {}


def px(emu):
    return int(round(emu / EMU_PER_CM * PX_PER_CM))


def font(size_pt, bold=False, italic=False):
    size = max(6, int(round(size_pt * 0.03528 * PX_PER_CM)))
    key = (size, bold, italic)
    if key not in _cache:
        _cache[key] = ImageFont.truetype("C:/Windows/Fonts/" + FONTS[(bool(bold), bool(italic))], size)
    return _cache[key]


def color_of(fmt_color, default=(0, 0, 0)):
    try:
        if fmt_color and fmt_color.type is not None:
            return tuple(fmt_color.rgb)
    except Exception:
        pass
    return default


def run_style(run, para_default=24):
    f = run.font
    size = f.size.pt if f.size else para_default
    return size, bool(f.bold), bool(f.italic), color_of(f.color)


def draw_text_frame(draw, tf, x, y, w, h, default_size=24):
    ml = px(tf.margin_left or 0) if tf.margin_left is not None else px(91440)
    mt = px(tf.margin_top or 0) if tf.margin_top is not None else px(45720)
    mr = px(tf.margin_right or 0) if tf.margin_right is not None else px(91440)
    cy = y + mt
    lines_out = []
    for p in tf.paragraphs:
        pPr = p._p.find(qn("a:pPr"))
        marl = px(int(pPr.get("marL", 0))) if pPr is not None else 0
        indent = px(int(pPr.get("indent", 0))) if pPr is not None else 0
        bullet = pPr is not None and pPr.find(qn("a:buChar")) is not None
        algn = pPr.get("algn", "l") if pPr is not None else "l"
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        runs = [(r.text, *run_style(r, default_size)) for r in p.runs]
        if not runs:
            cy += int(default_size * 0.03528 * PX_PER_CM * 1.2)
            continue
        size = max(r[1] for r in runs)
        lh = int(size * 0.03528 * PX_PER_CM * INTER_LINE * spacing)
        avail = (w - ml - mr - marl) / INTER_WIDTH
        # разбиваем на слова, сохраняя стиль
        words = []
        for text, sz, b, it, col in runs:
            for token in text.replace("\n", " ").split(" "):
                words.append((token, sz, b, it, col))
                words.append((" ", sz, b, it, col))
        line, lw, lines = [], 0, []
        for token, sz, b, it, col in words:
            tw = font(sz, b, it).getlength(token)
            if token != " " and lw + tw > avail and line:
                lines.append(line)
                line, lw = [], 0
            if token == " " and not line:
                continue
            line.append((token, sz, b, it, col))
            lw += tw
        if line:
            lines.append(line)
        for k, ln in enumerate(lines):
            total = sum(font(sz, b, it).getlength(t) for t, sz, b, it, col in ln)
            lx = x + ml + marl
            if algn == "ctr":
                lx = x + ml + (w - ml - mr - total) / 2
            elif algn == "r":
                lx = x + w - mr - total
            if bullet and k == 0:
                draw.text((x + ml + marl + indent, cy), "•", font=font(size), fill=(67, 75, 154))
            for t, sz, b, it, col in ln:
                draw.text((lx, cy + (lh - sz * 0.03528 * PX_PER_CM * 1.2) * 0.5), t, font=font(sz, b, it), fill=col)
                lx += font(sz, b, it).getlength(t)
            cy += lh
        sa = p.space_after.pt if p.space_after is not None else 0
        cy += int(sa * 0.03528 * PX_PER_CM)
    if cy > y + h + 2:
        lines_out.append(cy - (y + h))
    return lines_out


def main():
    prs = Presentation(SRC)
    W, H = px(prs.slide_width), px(prs.slide_height)
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    overflow = []
    for shape in prs.slides[0].shapes:
        x, y, w, h = px(shape.left), px(shape.top), px(shape.width), px(shape.height)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pic = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA").resize((max(w, 1), max(h, 1)))
            img.paste(pic, (x, y), pic)
            continue
        if shape.has_table if hasattr(shape, "has_table") else False:
            cy = y
            for r, row in enumerate(shape.table.rows):
                cx = x
                rh = px(row.height)
                for c, cell in enumerate(row.cells):
                    cw = px(shape.table.columns[c].width)
                    fill = color_of(cell.fill.fore_color, (255, 255, 255)) if cell.fill.type else (255, 255, 255)
                    draw.rectangle([cx, cy, cx + cw, cy + rh], fill=fill, outline=(200, 200, 210))
                    draw_text_frame(draw, cell.text_frame, cx, cy + rh * 0.2, cw, rh)
                    cx += cw
                cy += rh
            continue
        sp = shape._element.spPr if hasattr(shape._element, "spPr") else None
        geom = sp.find(qn("a:prstGeom")).get("prst") if sp is not None and sp.find(qn("a:prstGeom")) is not None else ""
        fill = None
        try:
            if shape.fill.type == 1:
                fill = color_of(shape.fill.fore_color, None)
        except Exception:
            pass
        line = None
        try:
            ln = sp.find(qn("a:ln")) if sp is not None else None
            if ln is not None and ln.find(qn("a:solidFill")) is not None:
                line = color_of(shape.line.color, (98, 102, 106))
        except Exception:
            pass
        if fill or line:
            box = [x, y, x + w, y + h]
            if geom == "ellipse":
                draw.ellipse(box, fill=fill, outline=line)
            elif geom == "roundRect":
                draw.rounded_rectangle(box, radius=int(min(w, h) * 0.025), fill=fill, outline=line, width=2)
            else:
                draw.rectangle(box, fill=fill, outline=line, width=2)
        if shape.has_text_frame and shape.text_frame.text.strip():
            over = draw_text_frame(draw, shape.text_frame, x, y, w, h)
            if over and over[0] > 4:
                overflow.append((shape.name, shape.text_frame.text[:40], round(over[0] / PX_PER_CM, 2)))
    img.save(OUT)
    print("превью:", OUT, img.size)
    for name, text, cm in overflow:
        print(f"  текст выходит за рамку на {cm} см: {name}: {text!r}")


if __name__ == "__main__":
    main()
